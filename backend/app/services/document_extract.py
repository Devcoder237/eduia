"""Extraction de texte des documents importés (PDF, Word, PowerPoint).

Limites assumées : texte sélectionnable uniquement (pas d'OCR : trop lourd pour
512 Mo de RAM). Protections : nombre de pages borné, taille décompressée des
.docx/.pptx bornée (anti « zip bomb »), texte tronqué à ``MAX_CHARS``.
Les fonctions sont synchrones : les appeler via ``anyio.to_thread.run_sync``.
"""

import io
import zipfile
from dataclasses import dataclass

from app.core.logging import get_logger

log = get_logger(__name__)

MAX_CHARS = 8000
MAX_PDF_PAGES = 40
MAX_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".pptx"})


@dataclass(frozen=True)
class ExtractionResult:
    """Résultat d'extraction (``error`` renseigné en cas d'échec)."""

    text: str
    truncated: bool
    char_count: int
    error: str | None


def _fail(message: str) -> ExtractionResult:
    return ExtractionResult(text="", truncated=False, char_count=0, error=message)


def get_extension(filename: str | None) -> str:
    """Retourne l'extension en minuscules (avec le point), ou ``""``."""
    if not filename or "." not in filename:
        return ""
    return "." + filename.rsplit(".", 1)[-1].lower()


def _check_zip_size(file_bytes: bytes) -> None:
    """Refuse les archives Office dont la taille décompressée est excessive."""
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
        total = sum(info.file_size for info in archive.infolist())
    if total > MAX_UNCOMPRESSED_BYTES:
        raise ValueError("document trop volumineux une fois décompressé")


def _extract_pdf(file_bytes: bytes) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages[:MAX_PDF_PAGES]:
            page_text = page.extract_text()
            if page_text:
                parts.append(page_text)
            if sum(len(p) for p in parts) > MAX_CHARS * 2:
                break
    return "\n\n".join(parts)


def _extract_docx(file_bytes: bytes) -> str:
    import docx

    _check_zip_size(file_bytes)
    document = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_pptx(file_bytes: bytes) -> str:
    from pptx import Presentation

    _check_zip_size(file_bytes)
    presentation = Presentation(io.BytesIO(file_bytes))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        slide_text = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if slide_text:
            parts.append(f"[Diapositive {index}]\n" + "\n".join(slide_text))
    return "\n\n".join(parts)


def extract_text(filename: str, file_bytes: bytes) -> ExtractionResult:
    """Extrait le texte d'un document sans jamais lever d'exception.

    Args:
        filename: Nom du fichier (sert à déduire le format).
        file_bytes: Contenu brut.

    Returns:
        Le résultat ; ``error`` contient un message utilisateur si échec.
    """
    ext = get_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        return _fail(
            f"Format {ext or 'inconnu'} non supporté (PDF, Word .docx, PowerPoint .pptx uniquement)."
        )

    extractors = {".pdf": _extract_pdf, ".docx": _extract_docx, ".pptx": _extract_pptx}
    try:
        raw_text = extractors[ext](file_bytes)
    except Exception as exc:  # noqa: BLE001 - les bibliothèques de parsing lèvent des types variés
        log.warning("document_extract_failed", ext=ext, error_type=exc.__class__.__name__)
        return _fail("Impossible de lire ce fichier. Il est peut-être corrompu, protégé ou trop volumineux.")

    raw_text = raw_text.strip()
    if not raw_text:
        return _fail("Aucun texte exploitable trouvé (le document est peut-être une image scannée).")

    return ExtractionResult(
        text=raw_text[:MAX_CHARS],
        truncated=len(raw_text) > MAX_CHARS,
        char_count=len(raw_text),
        error=None,
    )
